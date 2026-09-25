#!/usr/bin/env python3
# -*- coding: utf-8 -*-
import os
import zipfile
from openpyxl import Workbook, load_workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.properties import CalcProperties
from calculadora_sheet import build_calculadora
from docs_sheets import build_como_usar, build_exemplos
from math_engine import full_result
from taxas_sheet import build_taxas

OUT = "/workspace/precifica/Precifica-precificacao.xlsx"

def add_names(wb, pairs):
    for name, ref in pairs:
        wb.defined_names.add(DefinedName(name=name, attr_text=ref))

def main():
    os.makedirs("/workspace/precifica", exist_ok=True)
    ex1 = full_result("ML Classico", 50, 2, 0, "voce", 0.04, "Margem %", margem=0.20)
    ex2 = full_result("Amazon Individual", 100, 0, 20, "voce", 0.0, "Lucro R$", lucro=30)
    ex3 = full_result("ML Classico", 50, 2, 0, "voce", 0.04, "Margem %", margem=0.20, preco=100)

    wb = Workbook()
    wb.remove(wb.active)
    wb.calculation = CalcProperties(calcMode="auto", fullCalcOnLoad=True)
    wb.properties.creator = "Precifica"
    wb.properties.title = "Precifica — planilha de precificacao 2026"
    wb.properties.subject = "Calculadora de preco para marketplaces BR"
    wb.properties.description = "Estimativa. Confirme no Seller Center / Simulador oficial. Nao e conselho contabil."
    wb.properties.keywords = "precifica,precificacao,marketplace,mercado livre,shopee,amazon,magalu"

    names = build_taxas(wb)
    add_names(wb, names)
    build_calculadora(wb)
    build_como_usar(wb)
    build_exemplos(wb, ex1, ex2, ex3)

    order = ["Calculadora", "Taxas", "Como usar", "Exemplos"]
    for i, title in enumerate(order):
        wb.move_sheet(title, offset=i - wb.sheetnames.index(title))
    wb.active = 0
    wb.save(OUT)

    assert os.path.isfile(OUT)
    with zipfile.ZipFile(OUT, "r") as z:
        bad = z.testzip()
        assert bad is None, "zip corrompido: {0}".format(bad)
        names_in_zip = z.namelist()
        assert "[Content_Types].xml" in names_in_zip
        assert any(n.startswith("xl/worksheets/") for n in names_in_zip)

    wb2 = load_workbook(OUT)
    assert wb2.sheetnames == ["Calculadora", "Taxas", "Como usar", "Exemplos"]
    calc = wb2["Calculadora"]
    taxas = wb2["Taxas"]
    for addr in ["H7", "H8", "H9", "H10", "H11", "H12", "A18", "O8", "O20", "O23", "O38", "O40", "O43", "O55"]:
        val = calc[addr].value
        assert isinstance(val, str) and val.startswith("="), "{0} deveria ser formula, veio {1!r}".format(addr, val)
    assert calc["C7"].value == 50
    assert calc["C10"].value == "voce"
    assert calc["C15"].value == "ML Classico"
    assert taxas["C8"].value == 0.12
    assert taxas["C13"].value == 6.5
    assert calc["C7"].protection.locked is False
    assert calc["H7"].protection.locked is True
    assert calc.protection.sheet is True
    assert "Estimativa. Confirme no Seller Center" in str(calc["A3"].value)
    dn = set(wb2.defined_names.keys())
    for n in ["ML_Classico_pct", "SH_Fix4", "AMZ_IndExtra", "ListaMarketplace"]:
        assert n in dn, n

    size = os.path.getsize(OUT)
    from math_engine import brl, pct_br
    print("=" * 72)
    print("ARQUIVO")
    print("  path: {0}".format(OUT))
    print("  size: {0} bytes ({1:.1f} KB)".format(size, size / 1024.0))
    print("  sheets: {0}".format(wb2.sheetnames))
    print("  formulas conferidas: H7 H8 H9 H10 H11 H12 A18 O8 O20 O23 O38 O40 O43 O55")
    print("  zip/xlsx: valido")
    print()
    print("3 CENARIOS (Python, mesma matematica das formulas)")
    print("-" * 72)
    for i, r in enumerate((ex1, ex2, ex3), start=1):
        kind = "preco-alvo" if i < 3 else "reverso (S informado)"
        print("Cenario {0} — {1}".format(i, kind))
        print("  Marketplace : {0}".format(r["marketplace"]))
        print("  Entradas    : custo {0} | embalagem {1} | frete {2} ({3}) | imposto {4} | modo {5}".format(
            brl(r["custo"]), brl(r["embalagem"]), brl(r["frete"]), r["quem"], pct_br(r["imposto"]), r["modo"]))
        print("  BASE        : {0}".format(brl(r["base"])))
        print("  Preco S     : {0}".format(brl(r["S"])))
        print("  Taxa        : {0}   ({1} + fixa {2})".format(brl(r["taxa"]), pct_br(r["r"]), brl(r["k"])))
        print("  Imposto R$  : {0}".format(brl(r["imposto_rs"])))
        print("  Frete cons. : {0}".format(brl(r["frete_c"])))
        print("  Lucro liq.  : {0}".format(brl(r["lucro"])))
        print("  Margem ef.  : {0}".format(pct_br(r["margem_ef"])))
        print()
    print("=" * 72)
    return ex1, ex2, ex3, size

if __name__ == "__main__":
    main()
