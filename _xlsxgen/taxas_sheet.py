from openpyxl.styles import Alignment
from styles import (
    GOLD, MONEY, PCT, box_thin, center, fill_aux, fill_dark, fill_gold_pale,
    fill_mid, fill_range, fill_soft, fill_white, font_aux_h, font_body, font_h,
    font_label, font_muted, font_section, font_small, header_block, left_c,
    left_t, page_setup, put, style_merge, yellow_input,
)

def build_taxas(wb):
    ws = wb.create_sheet("Taxas", 1)
    ws.sheet_properties.tabColor = GOLD
    for col, w in {"A": 42, "B": 3, "C": 16, "D": 16, "E": 62, "F": 26, "G": 30}.items():
        ws.column_dimensions[col].width = w
    header_block(
        ws, "G",
        "Premissas editaveis — a Calculadora le estas celulas, nunca valores cravados nas formulas",
    )
    page_setup(ws, "A1:G42", freeze="A6")

    put(ws, "A5",
        "Edite so as celulas amarelas. Nao altere a grafia das listas a direita (os menus da Calculadora dependem delas).",
        font_muted, fill_gold_pale, left_c)
    ws.merge_cells("A5:E5")
    fill_range(ws, "A5:E5", fill_gold_pale)

    put(ws, "A6", "Parametro", font_h, fill_dark, center)
    put(ws, "B6", "", font_h, fill_dark, center)
    put(ws, "C6", "Valor", font_h, fill_dark, center)
    put(ws, "D6", "Unidade", font_h, fill_dark, center)
    put(ws, "E6", "Notas", font_h, fill_dark, left_c)

    rows = [
        (7, "sec", "Mercado Livre", None, "", ""),
        (8, "pct", "ML Classico %", 0.12, "%", "Comissao tipica do anuncio Classico. Varia por categoria."),
        (9, "pct", "ML Premium %", 0.17, "%", "Comissao tipica do anuncio Premium."),
        (10, "money", "ML sem taxa fixa a partir de", 79, "R$", "A partir deste preco a taxa fixa deixa de incidir (estimativa)."),
        (11, "money", "ML faixa baixa ate", 12.50, "R$", "Abaixo deste valor a fixa e um percentual do proprio preco."),
        (12, "pct", "ML fixa < 12,50 (fracao do preco)", 0.50, "% do preco", "50% do preco de venda, somado a comissao percentual."),
        (13, "money", "ML fixa 12,50-78,99", 6.50, "R$", "Estimativa da antiga taxa fixa. Confirme no simulador — em 2026 parte virou custo operacional por peso/cubagem."),
        (15, "sec", "Shopee", None, "", ""),
        (16, "money", "Shopee limite faixa 1 (abaixo de)", 80, "R$", "Primeiro corte de tabela."),
        (17, "pct", "Shopee < 80 %", 0.20, "%", "Comissao da faixa baixa."),
        (18, "money", "Shopee < 80 fixa", 4, "R$", "Taxa fixa da faixa baixa."),
        (19, "money", "Shopee limite faixa 2 (abaixo de)", 100, "R$", "Segundo corte."),
        (20, "pct", "Shopee 80-99,99 %", 0.14, "%", "Comissao da segunda faixa."),
        (21, "money", "Shopee 80-99,99 fixa", 16, "R$", "Taxa fixa da segunda faixa."),
        (22, "money", "Shopee limite faixa 3 (abaixo de)", 200, "R$", "Terceiro corte."),
        (23, "pct", "Shopee 100-199,99 %", 0.14, "%", "Comissao da terceira faixa."),
        (24, "money", "Shopee 100-199,99 fixa", 20, "R$", "Taxa fixa da terceira faixa."),
        (25, "pct", "Shopee >= 200 %", 0.14, "%", "Comissao da faixa alta."),
        (26, "money", "Shopee >= 200 fixa", 26, "R$", "Taxa fixa da faixa alta. O teto antigo de comissao caiu."),
        (28, "sec", "Amazon BR", None, "", ""),
        (29, "pct", "Amazon referral default", 0.12, "%", "Referral fee padrao. Categorias especificas diferem — edite se a sua for outra."),
        (30, "money", "Amazon Individual extra", 2, "R$/item", "Tarifa extra por item no plano Individual."),
        (31, "money", "Amazon Professional extra", 0, "R$/item", "Plano Professional: sem extra por item nesta premissa."),
        (32, "money", "Amazon min referral", 1, "R$", "Referral minimo por unidade. Fee = MAX(preco x %, minimo) + extra."),
        (34, "sec", "Magalu", None, "", ""),
        (35, "pct", "Magalu %", 0.16, "%", "Estimativa. Comissao real varia por categoria e acordo."),
        (37, "sec", "Venda direta", None, "", ""),
        (38, "pct", "Venda direta %", 0.0, "%", "Sem comissao de marketplace. Imposto e custos continuam valendo."),
    ]
    named = {
        8: "ML_Classico_pct", 9: "ML_Premium_pct", 10: "ML_SemFixaDe",
        11: "ML_FaixaBaixa", 12: "ML_FixaPctBaixa", 13: "ML_FixaMeio",
        16: "SH_Lim1", 17: "SH_Pct1", 18: "SH_Fix1",
        19: "SH_Lim2", 20: "SH_Pct2", 21: "SH_Fix2",
        22: "SH_Lim3", 23: "SH_Pct3", 24: "SH_Fix3",
        25: "SH_Pct4", 26: "SH_Fix4",
        29: "AMZ_Pct", 30: "AMZ_IndExtra", 31: "AMZ_ProExtra", 32: "AMZ_Min",
        35: "MAGALU_Pct", 38: "DIRETA_Pct",
    }
    for r, kind, label, val, unit, note in rows:
        ws.row_dimensions[r].height = 22 if kind != "sec" else 24
        if kind == "sec":
            style_merge(ws, "A{0}:E{0}".format(r), label, font_section, fill_mid,
                        Alignment(horizontal="left", vertical="center", indent=1))
            continue
        put(ws, "A{0}".format(r), label, font_label, fill_white, left_c, border=box_thin)
        put(ws, "B{0}".format(r), "", font_body, fill_white, left_c, border=box_thin)
        fmt = PCT if kind == "pct" else MONEY
        yellow_input(ws, "C{0}".format(r), val, fmt=fmt)
        put(ws, "D{0}".format(r), unit, font_small, fill_white, center, border=box_thin)
        put(ws, "E{0}".format(r), note, font_small, fill_white, left_c, border=box_thin)

    put(ws, "F6", "Listas da Calculadora", font_h, fill_dark, center)
    put(ws, "G6", "Nao altere a grafia", font_h, fill_dark, center)
    put(ws, "F7", "Marketplace", font_aux_h, fill_soft, left_c)
    markets = [
        "ML Classico", "ML Premium", "Shopee", "Amazon Individual",
        "Amazon Professional", "Magalu", "Venda direta",
    ]
    for i, name in enumerate(markets):
        put(ws, "F{0}".format(8 + i), name, font_body, fill_white, left_c, border=box_thin)
        put(ws, "G{0}".format(8 + i), "grafia usada no SWITCH da Calculadora", font_small, fill_aux, left_c)
    put(ws, "F16", "Quem paga o frete", font_aux_h, fill_soft, left_c)
    put(ws, "F17", "voce", font_body, fill_white, left_c, border=box_thin)
    put(ws, "F18", "cliente", font_body, fill_white, left_c, border=box_thin)
    put(ws, "G17", "entra no custo", font_small, fill_aux, left_c)
    put(ws, "G18", "fica fora do preco", font_small, fill_aux, left_c)
    put(ws, "F20", "Modo", font_aux_h, fill_soft, left_c)
    put(ws, "F21", "Margem %", font_body, fill_white, left_c, border=box_thin)
    put(ws, "F22", "Lucro R$", font_body, fill_white, left_c, border=box_thin)
    put(ws, "G21", "lucro = % do preco de venda", font_small, fill_aux, left_c)
    put(ws, "G22", "lucro = valor fixo em reais", font_small, fill_aux, left_c)

    put(ws, "A40",
        "Nota 2026: o Mercado Livre substituiu, em varios casos, a taxa fixa classica por custo operacional variavel (peso real vs. cubado). "
        "Os R$ 6,50 sao uma estimativa de faixa para a planilha nao quebrar. O valor oficial e o do simulador com a sua conta.",
        font_muted, fill_gold_pale, left_t)
    ws.merge_cells("A40:E42")
    fill_range(ws, "A40:E42", fill_gold_pale)

    names = [(name, "'Taxas'!$C${0}".format(row)) for row, name in named.items()]
    names.append(("ListaMarketplace", "'Taxas'!$F$8:$F$14"))
    names.append(("ListaFreteQuem", "'Taxas'!$F$17:$F$18"))
    names.append(("ListaModo", "'Taxas'!$F$21:$F$22"))
    return names
