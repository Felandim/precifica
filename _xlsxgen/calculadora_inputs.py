from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from formulas import faixa_of, k_of, price_fx, r_of
from styles import (
    GREEN_DARK, MONEY, PCT, NUM2, PWD, RED, box, box_thin, center, fill_aux,
    fill_dark, fill_gold_pale, fill_loss, fill_mid, fill_price, fill_range,
    fill_result, fill_soft, fill_white, font_aux, font_aux_h, font_body,
    font_h, font_label, font_muted, font_out, font_price, font_small,
    font_white, header_block, left_c, left_t, locked, money_out, page_setup,
    pct_out, put, right_c, style_merge, yellow_input, border_range,
)

def _inputs(ws):
    put(ws, "A5", "Amarelo = voce preenche  |  Verde = resultado (formula, bloqueado)", font_muted, fill_soft, left_c)
    ws.merge_cells("A5:D5")
    fill_range(ws, "A5:D5", fill_soft)
    put(ws, "F5", "Use o modo e o marketplace; a taxa NAO e digitada — vem da aba Taxas.", font_muted, fill_soft, left_c)
    ws.merge_cells("F5:L5")
    fill_range(ws, "F5:L5", fill_soft)
    style_merge(ws, "A6:D6", "ENTRADAS", font_h, fill_dark, Alignment(horizontal="left", vertical="center", indent=1))
    style_merge(ws, "F6:L6", "RESULTADOS", font_h, fill_dark, Alignment(horizontal="left", vertical="center", indent=1))
    inputs = [
        (7, "Custo do produto", 50, MONEY, "Quanto voce paga no item (sem embalagem)."),
        (8, "Embalagem", 2, MONEY, "Caixa, envelope, fita, etiqueta."),
        (9, "Frete", 0, MONEY, "Custo de envio. So entra no preco se voce pagar."),
        (10, "Quem paga o frete", "voce", None, "voce = soma no custo | cliente = fica de fora."),
        (11, "Imposto %", 0.04, PCT, "Sobre o preco de venda (MEI / Simples — confira seu regime)."),
        (12, "Modo", "Margem %", None, "Margem % ou Lucro R$ — so um dos dois e usado."),
        (13, "Margem %", 0.20, PCT, "Fatia do preco de venda que voce quer de lucro. Usada se o modo for Margem %."),
        (14, "Lucro R$", 30, MONEY, "Quanto quer ganhar por unidade. Usado se o modo for Lucro R$."),
        (15, "Marketplace", "ML Classico", None, "Define comissao e taxa fixa a partir da aba Taxas."),
    ]
    for r, label, val, fmt, hint in inputs:
        ws.row_dimensions[r].height = 22
        put(ws, "A{0}".format(r), label, font_label, fill_white, left_c, border=box_thin)
        put(ws, "B{0}".format(r), "", font_body, fill_white, left_c, border=box_thin)
        yellow_input(ws, "C{0}".format(r), val, fmt=fmt)
        put(ws, "D{0}".format(r), hint, font_small, fill_white, left_c, border=box_thin)
    border_range(ws, "A6:D15", box)

def _results(ws):
    ws.row_dimensions[7].height = 32
    put(ws, "F7", "Preco de venda", font_white, fill_price, left_c)
    put(ws, "G7", "", font_white, fill_price)
    money_out(ws, "H7", '=IFERROR($O$40,"Reduza taxas ou margem")', big=True)
    put(ws, "I7", "", font_white, fill_price)
    put(ws, "J7", "Quanto anunciar", font_small, fill_price, left_c)
    put(ws, "K7", "", font_white, fill_price)
    put(ws, "L7", "", font_white, fill_price)
    fill_range(ws, "F7:L7", fill_price)
    ws["H7"].font = font_price
    ws["H7"].fill = fill_price
    ws["H7"].alignment = center
    ws["J7"].font = Font(name="Calibri", size=9, italic=True, color="D8F3E8")
    rows = [
        (8, "Lucro liquido", "=$O$45", MONEY, "O que sobra depois de taxa, imposto, custo e frete."),
        (9, "Taxa em R$", "=$O$43", MONEY, "Comissao % + parcela fixa da faixa."),
        (10, "Imposto em R$", "=$O$44", MONEY, "Imposto % x preco de venda."),
        (11, "Frete considerado", "=$O$47", MONEY, "Zero se o cliente paga."),
        (12, "Margem efetiva", "=$O$46", PCT, "Lucro liquido / preco de venda."),
        (13, "Custo total", "=$O$8", MONEY, "Produto + embalagem + frete considerado."),
        (14, "Faixa de taxa aplicada", "=$O$48", None, "Qual regra da aba Taxas entrou no calculo."),
        (15, "Comissao % + fixa usadas", "=$O$41", PCT, "Percentual efetivo da faixa (a fixa em R$ ao lado)."),
    ]
    for r, label, formula, fmt, hint in rows:
        ws.row_dimensions[r].height = 22
        put(ws, "F{0}".format(r), label, font_label, fill_white, left_c, border=box_thin)
        put(ws, "G{0}".format(r), "", font_body, fill_white, left_c, border=box_thin)
        if fmt == PCT:
            pct_out(ws, "H{0}".format(r), formula)
        elif fmt == MONEY:
            money_out(ws, "H{0}".format(r), formula, big=False)
        else:
            put(ws, "H{0}".format(r), formula, font_out, fill_result, left_c, prot=locked, border=box_thin)
        put(ws, "J{0}".format(r), hint, font_small, fill_white, left_c)
        ws.merge_cells("J{0}:L{0}".format(r))
        fill_range(ws, "J{0}:L{0}".format(r), fill_white)
    put(ws, "I15", "=$O$42", font_out, fill_result, right_c, fmt=MONEY, prot=locked, border=box_thin)
    border_range(ws, "F6:L15", box)
    ws.conditional_formatting.add("H8", CellIsRule(operator="lessThan", formula=["0"], fill=fill_loss,
        font=Font(name="Calibri", size=12, bold=True, color=RED)))
    ws.conditional_formatting.add("H12", CellIsRule(operator="lessThan", formula=["0"], fill=fill_loss,
        font=Font(name="Calibri", size=12, bold=True, color=RED)))
    ws.conditional_formatting.add("C13", FormulaRule(formula=['$C$12="Lucro R$"'], fill=PatternFill("solid", fgColor="EEEEEE")))
    ws.conditional_formatting.add("C14", FormulaRule(formula=['$C$12="Margem %"'], fill=PatternFill("solid", fgColor="EEEEEE")))

def _breakdown_and_reverse(ws):
    style_merge(ws, "A17:L17", "BREAKDOWN — texto para colar nos grupos de sellers", font_h, fill_mid,
                Alignment(horizontal="left", vertical="center", indent=1))
    breakdown = (
        '="*Precifica — calculo de preco*"&CHAR(10)'
        '&"Marketplace: "&C15&CHAR(10)'
        '&"Custo do produto: "&TEXT(C7,"R$ #,##0.00")&CHAR(10)'
        '&"Embalagem: "&TEXT(C8,"R$ #,##0.00")&CHAR(10)'
        '&"Frete ("&IF(C10="voce","voce paga","cliente paga")&"): "&TEXT(C9,"R$ #,##0.00")&CHAR(10)'
        '&"Taxa: "&TEXT(O41,"0.00%")&" + fixa "&TEXT(O42,"R$ #,##0.00")&" = "&TEXT(O43,"R$ #,##0.00")&CHAR(10)'
        '&"Imposto: "&TEXT(O10,"0.00%")&" ("&TEXT(O44,"R$ #,##0.00")&")"&CHAR(10)'
        '&"*Preco de venda: "&TEXT(O40,"R$ #,##0.00")&"*"&CHAR(10)'
        '&"*Lucro liquido: "&TEXT(O45,"R$ #,##0.00")&"*"&CHAR(10)'
        '&"Margem efetiva: "&TEXT(O46,"0.00%")&CHAR(10)'
        '&"Faixa: "&O48&CHAR(10)&CHAR(10)'
        '&"Estimativa Precifica 2026. Confirme no Seller Center / Simulador oficial. Nao e conselho contabil."'
    )
    ws.merge_cells("A18:L21")
    put(ws, "A18", breakdown, Font(name="Calibri", size=10, color="1A1A1A"), fill_soft, left_t, prot=locked, border=box)
    fill_range(ws, "A18:L21", fill_soft)
    border_range(ws, "A18:L21", box)
    for r in (18, 19, 20, 21):
        ws.row_dimensions[r].height = 22

    style_merge(ws, "A23:L23",
                "SE VENDER A R$ X, QUANTO SOBRA?  —  nao precisa inverter a formula; informe o preco praticado",
                font_h, fill_dark, Alignment(horizontal="left", vertical="center", indent=1))
    put(ws, "A24", "Preco praticado (R$)", font_label, fill_white, left_c, border=box_thin)
    put(ws, "B24", "", font_body, fill_white, border=box_thin)
    yellow_input(ws, "C24", 100, fmt=MONEY)
    put(ws, "D24", "Preco que voce ja anuncia ou quer testar.", font_small, fill_white, left_c, border=box_thin)
    put(ws, "F24", "Lucro liquido nesse preco", font_label, fill_white, left_c, border=box_thin)
    put(ws, "G24", "", font_body, fill_white, border=box_thin)
    money_out(ws, "H24", "=$O$55")
    put(ws, "J24", "S - taxa - imposto - custo total", font_small, fill_white, left_c)
    ws.merge_cells("J24:L24")
    put(ws, "A25", "Custo total usado", font_label, fill_white, left_c, border=box_thin)
    put(ws, "B25", "", font_body, fill_white, border=box_thin)
    money_out(ws, "C25", "=$O$8")
    ws["C25"].protection = locked
    put(ws, "D25", "Mesmo custo da esquerda (nao edite aqui).", font_small, fill_white, left_c)
    put(ws, "F25", "Margem efetiva", font_label, fill_white, left_c, border=box_thin)
    put(ws, "G25", "", font_body, fill_white, border=box_thin)
    pct_out(ws, "H25", "=$O$56")
    put(ws, "J25", "Lucro / preco praticado", font_small, fill_white, left_c)
    ws.merge_cells("J25:L25")
    put(ws, "A26", "Faixa nesse preco", font_label, fill_white, left_c, border=box_thin)
    put(ws, "B26", "", font_body, fill_white, border=box_thin)
    ws.merge_cells("C26:D26")
    put(ws, "C26", "=$O$57", font_out, fill_result, left_c, prot=locked, border=box_thin)
    put(ws, "F26", "Taxa em R$", font_label, fill_white, left_c, border=box_thin)
    put(ws, "G26", "", font_body, fill_white, border=box_thin)
    money_out(ws, "H26", "=$O$53")
    put(ws, "J26", "Comissao da faixa do preco praticado", font_small, fill_white, left_c)
    ws.merge_cells("J26:L26")
    put(ws, "F27", "Imposto em R$", font_label, fill_white, left_c, border=box_thin)
    put(ws, "G27", "", font_body, fill_white, border=box_thin)
    money_out(ws, "H27", "=$O$54")
    put(ws, "J27", "Imposto % x preco praticado", font_small, fill_white, left_c)
    ws.merge_cells("J27:L27")
    ws.conditional_formatting.add("H24", CellIsRule(operator="lessThan", formula=["0"], fill=fill_loss,
        font=Font(name="Calibri", size=12, bold=True, color=RED)))
    border_range(ws, "A23:L27", box)
    put(ws, "A29",
        "Celulas auxiliares (colunas N-P): preco-teste, um candidato por faixa, e o preco final consistente com a propria faixa. "
        "Nao edite. Assim a planilha nao oscila no corte de R$ 79. Impressao omite essa area. Senha da protecao: precifica (celulas amarelas continuam liberadas).",
        font_muted, fill_aux, left_t)
    ws.merge_cells("A29:L31")
    fill_range(ws, "A29:L31", fill_aux)

def _validations_and_protect(ws):
    dv_mkt = DataValidation(type="list", formula1="=ListaMarketplace", allow_blank=False, showDropDown=False)
    dv_mkt.error = "Escolha um marketplace da lista."
    dv_mkt.errorTitle = "Marketplace"
    dv_mkt.add("C15")
    ws.add_data_validation(dv_mkt)
    dv_quem = DataValidation(type="list", formula1="=ListaFreteQuem", allow_blank=False, showDropDown=False)
    dv_quem.error = "Escolha voce ou cliente."
    dv_quem.errorTitle = "Frete"
    dv_quem.add("C10")
    ws.add_data_validation(dv_quem)
    dv_modo = DataValidation(type="list", formula1="=ListaModo", allow_blank=False, showDropDown=False)
    dv_modo.error = "Escolha Margem % ou Lucro R$."
    dv_modo.errorTitle = "Modo"
    dv_modo.add("C12")
    ws.add_data_validation(dv_modo)
    ws.protection.password = PWD
    ws.protection.sheet = True
    ws.protection.enable()
    ws.protection.selectLockedCells = True
    ws.protection.selectUnlockedCells = True
