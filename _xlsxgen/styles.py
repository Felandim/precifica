# styles and small helpers
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import range_boundaries
from openpyxl.worksheet.page import PageMargins

GREEN_DARK = "0F4C3A"
GREEN_MID = "1A6B54"
GREEN_SOFT = "E6F4EF"
GREEN_RESULT = "D8F3E8"
GOLD = "C9A227"
GOLD_PALE = "FFF8E1"
YELLOW = "FFF2CC"
WHITE = "FFFFFF"
INK = "1A1A1A"
MUTED = "5C5C5C"
LINE = "C5D5CE"
AUX = "F4F6F5"
RED = "B42318"
LOSS_BG = "FDECEC"
LINK = "0563C1"
PWD = "precifica"
MONEY = r'"R$" #,##0.00'
PCT = "0.00%"
NUM2 = "#,##0.00"

thin = Side(style="thin", color=LINE)
med = Side(style="medium", color=GREEN_DARK)
gold_med = Side(style="medium", color=GOLD)
font_title = Font(name="Calibri", size=22, bold=True, color=WHITE)
font_sub = Font(name="Calibri", size=11, italic=True, color=WHITE)
font_h = Font(name="Calibri", size=12, bold=True, color=WHITE)
font_label = Font(name="Calibri", size=11, bold=True, color=INK)
font_body = Font(name="Calibri", size=11, color=INK)
font_muted = Font(name="Calibri", size=9, italic=True, color=MUTED)
font_small = Font(name="Calibri", size=9, color=MUTED)
font_input = Font(name="Calibri", size=12, bold=True, color=INK)
font_price = Font(name="Calibri", size=20, bold=True, color=WHITE)
font_out = Font(name="Calibri", size=12, bold=True, color=GREEN_DARK)
font_link = Font(name="Calibri", size=10, underline="single", color=LINK)
font_disc = Font(name="Calibri", size=10, italic=True, color="7A5C00")
font_aux = Font(name="Calibri", size=9, color=MUTED)
font_aux_h = Font(name="Calibri", size=10, bold=True, color=GREEN_DARK)
font_white = Font(name="Calibri", size=11, bold=True, color=WHITE)
font_section = Font(name="Calibri", size=11, bold=True, color=WHITE)
fill_dark = PatternFill("solid", fgColor=GREEN_DARK)
fill_mid = PatternFill("solid", fgColor=GREEN_MID)
fill_soft = PatternFill("solid", fgColor=GREEN_SOFT)
fill_result = PatternFill("solid", fgColor=GREEN_RESULT)
fill_yellow = PatternFill("solid", fgColor=YELLOW)
fill_gold_pale = PatternFill("solid", fgColor=GOLD_PALE)
fill_white = PatternFill("solid", fgColor=WHITE)
fill_aux = PatternFill("solid", fgColor=AUX)
fill_price = PatternFill("solid", fgColor=GREEN_DARK)
fill_loss = PatternFill("solid", fgColor=LOSS_BG)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
left_c = Alignment(horizontal="left", vertical="center", wrap_text=True)
left_t = Alignment(horizontal="left", vertical="top", wrap_text=True)
right_c = Alignment(horizontal="right", vertical="center")
unlocked = Protection(locked=False)
locked = Protection(locked=True)
box = Border(left=med, right=med, top=med, bottom=med)
box_thin = Border(left=thin, right=thin, top=thin, bottom=thin)
box_gold = Border(left=gold_med, right=gold_med, top=gold_med, bottom=gold_med)

def fill_range(ws, rng, fill):
    min_col, min_row, max_col, max_row = range_boundaries(rng)
    for r in range(min_row, max_row + 1):
        for c in range(min_col, max_col + 1):
            ws.cell(r, c).fill = fill

def border_range(ws, rng, border):
    min_col, min_row, max_col, max_row = range_boundaries(rng)
    for r in range(min_row, max_row + 1):
        for c in range(min_col, max_col + 1):
            ws.cell(r, c).border = border

def style_merge(ws, rng, value, font, fill, alignment, border=None):
    ws.merge_cells(rng)
    cell = ws[rng.split(":")[0]]
    cell.value = value
    cell.font = font
    cell.fill = fill
    cell.alignment = alignment
    cell.protection = locked
    fill_range(ws, rng, fill)
    if border:
        border_range(ws, rng, border)
    return cell

def put(ws, addr, value, font=font_body, fill=None, alignment=left_c, fmt=None, prot=locked, border=None):
    cell = ws[addr]
    cell.value = value
    cell.font = font
    cell.alignment = alignment
    cell.protection = prot
    if fill is not None:
        cell.fill = fill
    if fmt:
        cell.number_format = fmt
    if border is not None:
        cell.border = border
    return cell

def yellow_input(ws, addr, value, fmt=None):
    return put(ws, addr, value, font=font_input, fill=fill_yellow, alignment=center, fmt=fmt, prot=unlocked, border=box_gold)

def money_out(ws, addr, formula, big=False):
    return put(ws, addr, formula, font=font_price if big else font_out, fill=fill_price if big else fill_result,
               alignment=center if big else right_c, fmt=MONEY, prot=locked, border=box if big else box_thin)

def pct_out(ws, addr, formula):
    return put(ws, addr, formula, font=font_out, fill=fill_result, alignment=right_c, fmt=PCT, prot=locked, border=box_thin)

def page_setup(ws, print_area, landscape=True, freeze="A6"):
    ws.freeze_panes = freeze
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_setup.horizontalCentered = True
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.45, right=0.45, top=0.5, bottom=0.5, header=0.25, footer=0.25)
    ws.print_options.horizontalCentered = True
    ws.print_area = print_area
    ws.oddHeader.left.text = "Precifica"
    ws.oddFooter.left.text = "Estimativa. Confirme no Seller Center. Nao e conselho contabil."
    ws.oddFooter.right.text = "Pagina &P de &N"
    ws.sheet_view.showGridLines = False

def header_block(ws, last_col, subtitle):
    ws.row_dimensions[1].height = 34
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 20
    ws.row_dimensions[4].height = 18
    style_merge(ws, f"A1:{last_col}1", "Precifica — planilha de precificacao 2026", font_title, fill_dark,
                Alignment(horizontal="left", vertical="center", indent=1))
    style_merge(ws, f"A2:{last_col}2", subtitle, font_sub, fill_mid, Alignment(horizontal="left", vertical="center", indent=1))
    style_merge(ws, f"A3:{last_col}3",
                "Estimativa. Confirme no Seller Center / Simulador oficial. Nao e conselho contabil.",
                font_disc, fill_gold_pale, Alignment(horizontal="left", vertical="center", indent=1))
    cell = style_merge(ws, f"A4:{last_col}4",
                       "Simulador oficial de custos do Mercado Livre -> https://www.mercadolivre.com.br/simulador-de-custos",
                       font_link, fill_white, Alignment(horizontal="left", vertical="center", indent=1))
    cell.hyperlink = "https://www.mercadolivre.com.br/simulador-de-custos"
